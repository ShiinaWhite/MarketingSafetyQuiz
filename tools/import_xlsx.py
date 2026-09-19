# -*- coding: utf-8 -*-
"""
tools/import_xlsx.py —— 题库导入脚本
把本地的安规题库 Excel（A列序号 B列专业 C列题型 D列题目 E列选项 F列答案）转换为
www/data/questions.json 与 www/data/questions.js，刷题程序只读生成结果，不依赖 Excel。

本仓库不附带任何题库数据，请使用你自己的题库 Excel 运行本脚本。

用法:
    python -X utf8 tools/import_xlsx.py <题库.xlsx>
不带参数时，自动在 当前目录 与 仓库根目录 中查找第一个 .xlsx 文件。
"""
import datetime
import json
import random
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

REPO_ROOT = Path(__file__).resolve().parent.parent
WEB_DATA_DIR = REPO_ROOT / "www" / "data"

# 题型 -> 期望题数（与用户核对一致）
EXPECTED = {"单选题": 160, "多选题": 110, "判断题": 122}
LETTERS = "ABCDEF"

# 全角字母/句点统一转半角，避免 "Ａ．" 之类写法漏切
FW_MAP = str.maketrans({"Ａ": "A", "Ｂ": "B", "Ｃ": "C", "Ｄ": "D",
                        "Ｅ": "E", "Ｆ": "F", "．": ".", "　": " "})
LABEL_RE = re.compile(r"([A-F])\s*[\.、]")  # 选项标签边界：字母 + 点/顿号


def norm(s):
    return str(s).translate(FW_MAP)


def clean_body(t):
    """去掉选项正文两端的分隔符；部分题库正文以 '- ' 开头，同样去掉便于阅读。"""
    t = t.strip().strip(",，;；").strip()
    if t.startswith("-"):
        t = t[1:].lstrip()
    return t.strip()


def _split_by_labels(s):
    """按标签边界切成 [(label, body), ...]。"""
    marks = [(m.start(), m.group(1)) for m in LABEL_RE.finditer(s)]
    out = []
    for i, (pos, ch) in enumerate(marks):
        start = pos + 1
        m = re.match(r"\s*[\.、]\s*", s[start:])
        start += m.end() if m else 0
        end = marks[i + 1][0] if i + 1 < len(marks) else len(s)
        out.append((ch, clean_body(s[start:end])))
    return out


def parse_options(raw):
    """E列 -> [(label, body), ...]。

    只保留从 A 开始、按字母连续递进的标签，防止选项正文里出现 "B." 之类的
    假标签把切分打断（题干正文可能含字母+点号，但 E 列只有选项文本）。
    """
    s = norm(raw).replace("\r", "").strip()
    if not s:
        return []
    picked, expect = [], 0
    for m in LABEL_RE.finditer(s):
        if m.group(1) == LETTERS[expect]:
            picked.append((m.start(), m.group(1)))
            expect += 1
    if not picked or picked[0][1] != "A":
        return []
    result = []
    for i, (pos, ch) in enumerate(picked):
        start = pos + 1
        m = re.match(r"\s*[\.、]\s*", s[start:])
        start += m.end() if m else 0
        end = picked[i + 1][0] if i + 1 < len(picked) else len(s)
        result.append((ch, clean_body(s[start:end])))
    return result


def parse_answer(raw, bodies):
    """F列 -> 正确选项下标列表。

    答案列是 "C. -移动, D. -拆除" 这类完整文本：按标签切分后，
    仅当“正文为空（纯字母答案）”或“正文能在 E 列选项中找到”才认为是真答案标签，
    避免答案正文里偶然出现的字母+点号被误判。
    """
    s = norm(raw).replace("\r", "").strip()
    if not s:
        return []
    body_set = {b for _, b in bodies if b}
    labels = []
    for ch, body in _split_by_labels(s):
        if body == "" or body in body_set:
            labels.append(ch)
    if not labels:  # 兜底：纯字母答案，如 "ABD"
        labels = [c for c in s.upper() if c in LETTERS]
    n = len(bodies)
    return sorted({LETTERS.index(c) for c in labels if LETTERS.index(c) < n})


def detect_type(t_name):
    if "单选" in t_name:
        return "single"
    if "多选" in t_name:
        return "multi"
    if "判断" in t_name:
        return "judge"
    return None


def find_default_xlsx():
    cands = []
    for d in (Path.cwd(), REPO_ROOT):
        cands.extend(sorted(d.glob("*.xlsx")))
    return cands[0] if cands else None


def main():
    if len(sys.argv) > 1:
        xlsx = Path(sys.argv[1])
    else:
        xlsx = find_default_xlsx()
        if xlsx is None:
            print("未找到题库 Excel。用法: python -X utf8 tools/import_xlsx.py <题库.xlsx>")
            sys.exit(2)
    print(f"读取: {xlsx}")
    wb = load_workbook(str(xlsx), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    # 跳过表头行
    start = 0
    for i, r in enumerate(rows[:3]):
        cells = [str(c or "") for c in r]
        if any(("题型" in c) or ("题目" in c) for c in cells):
            start = i + 1
            break

    questions, errors, warnings = [], [], []
    type_counts, seen_ids, dup_stems = Counter(), set(), Counter()

    for rno, r in enumerate(rows[start:], start=start + 1):
        vals = (list(r) + [None] * 6)[:6]
        seq_raw, major, t_raw, stem_raw, opts_raw, ans_raw = vals
        if all(v is None or str(v).strip() == "" for v in vals):
            continue

        t_name = str(t_raw or "").strip()
        t = detect_type(t_name)
        if t is None:
            errors.append(f"第{rno}行: 无法识别题型 {t_name!r}")
            continue

        stem = norm(stem_raw).replace("\r", "\n").strip()
        if not stem:
            errors.append(f"第{rno}行: 题干为空")
            continue
        dup_stems[stem] += 1

        bodies = parse_options(opts_raw)
        if t == "judge" and len(bodies) < 2:
            bodies = [("A", "正确"), ("B", "错误")]  # 判断题兜底
        if len(bodies) < 2:
            errors.append(f"第{rno}行: 选项解析失败: {str(opts_raw)[:60]!r}")
            continue
        # 选项正文里若混有假标签，解析出的标签数会少于全文标签数 -> 提示人工核查
        naive = len(LABEL_RE.findall(norm(opts_raw)))
        if naive != len(bodies):
            warnings.append(f"第{rno}行: E列出现 {naive} 个标签但只解析出 {len(bodies)} 个选项")

        answer = parse_answer(ans_raw, bodies)
        if not answer:
            errors.append(f"第{rno}行: 答案解析失败: {str(ans_raw)[:60]!r}")
            continue
        if max(answer) >= len(bodies):
            errors.append(f"第{rno}行: 答案 {answer} 超出选项数 {len(bodies)}")
            continue
        if t == "single" and len(answer) != 1:
            errors.append(f"第{rno}行: 单选题答案数={len(answer)}")
            continue
        if t == "multi" and len(answer) < 2:
            errors.append(f"第{rno}行: 多选题答案数={len(answer)}")
            continue

        try:
            seq = int(str(seq_raw).strip())
        except (TypeError, ValueError):
            seq = rno
        if seq in seen_ids:
            warnings.append(f"第{rno}行: 序号 {seq} 重复，改用行号 {rno} 作为 id")
            seq = rno
        seen_ids.add(seq)

        questions.append({
            "id": seq,
            "row": rno,
            "major": str(major or "").strip(),
            "type": t,
            "type_name": t_name,
            "stem": stem,
            "options": [b for _, b in bodies],
            "answer": answer,
            "answer_text": "".join(LETTERS[i] for i in answer),
        })
        type_counts[t_name] += 1

    # ---------- 校验报告 ----------
    total = len(questions)
    print(f"\n共导入 {total} 题")
    ok = True
    for t_name, expect in EXPECTED.items():
        got = type_counts.get(t_name, 0)
        mark = "OK " if got == expect else "FAIL"
        if got != expect:
            ok = False
        print(f"  [{mark}] {t_name}: {got} (期望 {expect})")

    opt_dist = Counter(len(q["options"]) for q in questions)
    print(f"选项数分布: {dict(sorted(opt_dist.items()))}")
    multi_ans = Counter(len(q["answer"]) for q in questions if q["type"] == "multi")
    print(f"多选答案个数分布: {dict(sorted(multi_ans.items()))}")
    same_stem = sum(c - 1 for c in dup_stems.values() if c > 1)
    print(f"题干重复(保留为独立题, 不去重): {same_stem} 组")

    if warnings:
        print(f"\n警告 {len(warnings)} 条:")
        for w in warnings[:20]:
            print("  !", w)
    if errors:
        print(f"\n错误 {len(errors)} 条:")
        for e in errors[:30]:
            print("  X", e)
        ok = False

    # ---------- 抽样打印 ----------
    rng = random.Random(2026)
    print("\n----- 抽样 -----")
    for t in ("single", "multi", "judge"):
        pool = [q for q in questions if q["type"] == t]
        if not pool:
            continue
        sample = [pool[0], pool[len(pool) // 2], pool[-1]] + rng.sample(pool, min(2, len(pool)))
        seen = set()
        for q in sample:
            if q["id"] in seen:
                continue
            seen.add(q["id"])
            print(f"\n[{q['type_name']} id={q['id']} {q['major']}] {q['stem'][:50]}")
            for i, opt in enumerate(q["options"]):
                mark = " <== 正确" if i in q["answer"] else ""
                print(f"   {LETTERS[i]}. {opt[:40]}{mark}")
            print(f"   答案: {q['answer_text']}")

    if not ok:
        print("\n校验未通过，不写出 questions.json")
        sys.exit(1)

    data = {
        "source": xlsx.name,
        "generated_at": datetime.datetime.now().isoformat(timespec="seconds"),
        "total": total,
        "type_counts": dict(type_counts),
        "questions": questions,
    }
    payload = json.dumps(data, ensure_ascii=False, indent=1)
    WEB_DATA_DIR.mkdir(parents=True, exist_ok=True)
    (WEB_DATA_DIR / "questions.json").write_text(payload, encoding="utf-8")
    # 内联副本：双击 index.html（file://）或任何无法 fetch 的环境也能零请求加载
    (WEB_DATA_DIR / "questions.js").write_text(
        "/* 由 tools/import_xlsx.py 自动生成，勿手改；请勿把题库数据提交到公共仓库 */\nwindow.QUESTIONS_DATA = " + payload + ";\n",
        encoding="utf-8")
    print(f"\n已写出: {WEB_DATA_DIR / 'questions.json'} ({(WEB_DATA_DIR / 'questions.json').stat().st_size} 字节)")
    print(f"已写出: {WEB_DATA_DIR / 'questions.js'}")


if __name__ == "__main__":
    main()
